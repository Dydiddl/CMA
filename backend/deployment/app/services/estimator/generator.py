from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import PatternFill, Border, Side, Alignment, Font
from pathlib import Path

class EstimatorGenerator:
    """표준공사내역서 자동 생성 클래스"""
    
    def __init__(self):
        self.template_dir = Path(__file__).parent / 'templates'
        self.default_template = self.template_dir / 'standard_template.xlsx'
        
        # 스타일 정의
        self.header_fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        self.header_font = Font(bold=True)
        self.center_alignment = Alignment(horizontal='center', vertical='center')
    
    def generate_estimate(self, 
                         work_items: List[Dict[str, Any]], 
                         output_path: str,
                         template_path: str = None) -> str:
        """
        공사내역서 생성
        
        Args:
            work_items (List[Dict[str, Any]]): 공사 항목 목록
            output_path (str): 출력 파일 경로
            template_path (str, optional): 템플릿 파일 경로
            
        Returns:
            str: 생성된 파일 경로
        """
        template = template_path or str(self.default_template)
        
        if not Path(template).exists():
            raise FileNotFoundError(f"템플릿 파일을 찾을 수 없습니다: {template}")
            
        wb = openpyxl.load_workbook(template)
        ws = wb.active
        
        # 데이터 입력
        self._fill_work_items(ws, work_items)
        
        # 스타일 적용
        self._apply_styles(ws)
        
        # 파일 저장
        wb.save(output_path)
        return output_path
    
    def _fill_work_items(self, worksheet, work_items: List[Dict[str, Any]]):
        """
        공사 항목 데이터를 워크시트에 입력
        
        Args:
            worksheet: 엑셀 워크시트
            work_items (List[Dict[str, Any]]): 공사 항목 목록
        """
        # 헤더 행 추가
        headers = ['코드', '공사항목', '수량', '단위', '금액']
        for col, header in enumerate(headers, 1):
            cell = worksheet.cell(row=1, column=col)
            cell.value = header
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.border = self.border
            cell.alignment = self.center_alignment
        
        # 데이터 입력
        for row, item in enumerate(work_items, 2):
            worksheet.cell(row=row, column=1, value=item['code'])
            worksheet.cell(row=row, column=2, value=item['name'])
            worksheet.cell(row=row, column=3, value=item['quantity'])
            worksheet.cell(row=row, column=4, value=item['unit'])
            worksheet.cell(row=row, column=5, value=item['amount'])
    
    def _apply_styles(self, worksheet):
        """
        워크시트에 스타일 적용
        
        Args:
            worksheet: 엑셀 워크시트
        """
        # 열 너비 조정
        worksheet.column_dimensions['A'].width = 15  # 코드
        worksheet.column_dimensions['B'].width = 40  # 공사항목
        worksheet.column_dimensions['C'].width = 15  # 수량
        worksheet.column_dimensions['D'].width = 10  # 단위
        worksheet.column_dimensions['E'].width = 20  # 금액
        
        # 모든 셀에 테두리 적용
        for row in worksheet.iter_rows():
            for cell in row:
                cell.border = self.border
                cell.alignment = self.center_alignment
    
    def create_template(self, output_path: str):
        """
        새로운 템플릿 파일 생성
        
        Args:
            output_path (str): 출력 파일 경로
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        
        # 기본 스타일 적용
        self._apply_styles(ws)
        
        # 파일 저장
        wb.save(output_path)
