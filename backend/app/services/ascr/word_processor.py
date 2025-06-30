#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ASCR Word 문서 처리 모듈
Word 문서(.docx) 분석 및 구조화
"""

import logging
from pathlib import Path
from typing import Dict, List, Any, Optional, Tuple
from dataclasses import dataclass
import json
import re
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.shared import OxmlElement, qn

logger = logging.getLogger(__name__)

@dataclass
class WordDocumentStructure:
    """Word 문서 구조 데이터 클래스"""
    title: str
    sections: List[Dict[str, Any]]
    tables: List[Dict[str, Any]]
    images: List[Dict[str, Any]]
    total_paragraphs: int
    total_tables: int
    total_images: int

@dataclass
class WordParagraph:
    """Word 단락 데이터 클래스"""
    text: str
    level: int
    style: str
    font_size: Optional[float]
    is_bold: bool
    is_italic: bool
    alignment: str

class WordDocumentProcessor:
    """Word 문서 처리 클래스"""
    
    def __init__(self):
        self.supported_extensions = ['.docx']
        self.construction_keywords = {
            "공통부문": ["공통", "일반", "기본", "표준", "규정"],
            "토목부문": ["토목", "도로", "교량", "터널", "댐", "제방"],
            "건축부문": ["건축", "건물", "시설", "구조", "마감"],
            "기계설비부문": ["기계", "설비", "전기", "통신", "소방"],
            "유지관리부문": ["유지", "관리", "보수", "점검", "정비"]
        }
    
    def can_process(self, file_path: Path) -> bool:
        """파일 처리 가능 여부 확인"""
        return file_path.suffix.lower() in self.supported_extensions
    
    def extract_text_from_word(self, file_path: Path) -> str:
        """Word 문서에서 텍스트 추출"""
        try:
            doc = Document(file_path)
            text_content = []
            
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content.append(paragraph.text)
            
            return '\n'.join(text_content)
            
        except Exception as e:
            logger.error(f"Word 문서 텍스트 추출 실패: {e}")
            raise
    
    def analyze_document_structure(self, file_path: Path) -> WordDocumentStructure:
        """Word 문서 구조 분석"""
        try:
            doc = Document(file_path)
            
            # 제목 추출
            title = self._extract_title(doc)
            
            # 섹션 분석
            sections = self._analyze_sections(doc)
            
            # 테이블 분석
            tables = self._analyze_tables(doc)
            
            # 이미지 분석
            images = self._analyze_images(doc)
            
            return WordDocumentStructure(
                title=title,
                sections=sections,
                tables=tables,
                images=images,
                total_paragraphs=len(doc.paragraphs),
                total_tables=len(doc.tables),
                total_images=len(images)
            )
            
        except Exception as e:
            logger.error(f"Word 문서 구조 분석 실패: {e}")
            raise
    
    def _extract_title(self, doc: Document) -> str:
        """문서 제목 추출"""
        # 첫 번째 제목 스타일 단락 찾기
        for paragraph in doc.paragraphs:
            if paragraph.style.name.startswith('Heading'):
                return paragraph.text.strip()
        
        # 제목이 없으면 첫 번째 단락 사용
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                return paragraph.text.strip()[:100]  # 100자로 제한
        
        return "제목 없음"
    
    def _analyze_sections(self, doc: Document) -> List[Dict[str, Any]]:
        """섹션 분석"""
        sections = []
        current_section = None
        
        for i, paragraph in enumerate(doc.paragraphs):
            text = paragraph.text.strip()
            if not text:
                continue
            
            # 제목 스타일 확인
            if paragraph.style.name.startswith('Heading'):
                # 이전 섹션 저장
                if current_section:
                    sections.append(current_section)
                
                # 새 섹션 시작
                level = int(paragraph.style.name.replace('Heading ', ''))
                current_section = {
                    "title": text,
                    "level": level,
                    "paragraph_number": i + 1,
                    "content": [],
                    "category": self._classify_section(text)
                }
            elif current_section:
                # 현재 섹션에 내용 추가
                current_section["content"].append({
                    "text": text,
                    "style": paragraph.style.name,
                    "paragraph_number": i + 1
                })
        
        # 마지막 섹션 추가
        if current_section:
            sections.append(current_section)
        
        return sections
    
    def _analyze_tables(self, doc: Document) -> List[Dict[str, Any]]:
        """테이블 분석"""
        tables = []
        
        for i, table in enumerate(doc.tables):
            table_data = {
                "table_number": i + 1,
                "rows": len(table.rows),
                "columns": len(table.columns),
                "data": []
            }
            
            # 테이블 데이터 추출
            for row_idx, row in enumerate(table.rows):
                row_data = []
                for cell_idx, cell in enumerate(row.cells):
                    row_data.append({
                        "text": cell.text.strip(),
                        "row": row_idx + 1,
                        "column": cell_idx + 1
                    })
                table_data["data"].append(row_data)
            
            # 테이블 유형 분류
            table_data["type"] = self._classify_table(table_data)
            tables.append(table_data)
        
        return tables
    
    def _analyze_images(self, doc: Document) -> List[Dict[str, Any]]:
        """이미지 분석"""
        images = []
        
        # Word 문서의 이미지 요소 찾기
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref:
                images.append({
                    "type": "image",
                    "target": rel.target_ref,
                    "relationship_id": rel.rId
                })
        
        return images
    
    def _classify_section(self, text: str) -> str:
        """섹션 분류"""
        text_lower = text.lower()
        
        for category, keywords in self.construction_keywords.items():
            for keyword in keywords:
                if keyword in text_lower:
                    return category
        
        return "기타"
    
    def _classify_table(self, table_data: Dict[str, Any]) -> str:
        """테이블 유형 분류"""
        if not table_data["data"]:
            return "빈 테이블"
        
        # 첫 번째 행의 내용으로 테이블 유형 판단
        first_row = table_data["data"][0]
        first_row_text = " ".join([cell["text"] for cell in first_row])
        
        if any(keyword in first_row_text for keyword in ["단가", "금액", "비용"]):
            return "단가표"
        elif any(keyword in first_row_text for keyword in ["수량", "개수", "면적"]):
            return "수량표"
        elif any(keyword in first_row_text for keyword in ["공종", "세부항목", "내역"]):
            return "내역서"
        else:
            return "일반표"
    
    def extract_structured_data(self, file_path: Path) -> Dict[str, Any]:
        """구조화된 데이터 추출"""
        try:
            structure = self.analyze_document_structure(file_path)
            text_content = self.extract_text_from_word(file_path)
            
            return {
                "file_path": str(file_path),
                "file_type": "word",
                "title": structure.title,
                "sections": structure.sections,
                "tables": structure.tables,
                "images": structure.images,
                "text_content": text_content,
                "statistics": {
                    "total_paragraphs": structure.total_paragraphs,
                    "total_tables": structure.total_tables,
                    "total_images": structure.total_images,
                    "total_sections": len(structure.sections)
                },
                "categories": self._extract_categories(structure.sections)
            }
            
        except Exception as e:
            logger.error(f"Word 문서 구조화 데이터 추출 실패: {e}")
            raise
    
    def _extract_categories(self, sections: List[Dict[str, Any]]) -> Dict[str, int]:
        """카테고리별 통계 추출"""
        categories = {}
        for section in sections:
            category = section.get("category", "기타")
            categories[category] = categories.get(category, 0) + 1
        return categories
    
    def create_word_template(self, template_data: Dict[str, Any], output_path: Path) -> Path:
        """Word 템플릿 생성"""
        try:
            doc = Document()
            
            # 제목 추가
            title = doc.add_heading(template_data.get("title", "건설 내역서"), 0)
            
            # 섹션 추가
            for section in template_data.get("sections", []):
                # 섹션 제목
                heading = doc.add_heading(section["title"], section.get("level", 1))
                
                # 섹션 내용
                for content in section.get("content", []):
                    paragraph = doc.add_paragraph(content["text"])
                    
                    # 스타일 적용
                    if content.get("is_bold"):
                        paragraph.runs[0].bold = True
                    if content.get("is_italic"):
                        paragraph.runs[0].italic = True
            
            # 테이블 추가
            for table_data in template_data.get("tables", []):
                if table_data["data"]:
                    rows = len(table_data["data"])
                    cols = len(table_data["data"][0]) if table_data["data"] else 1
                    
                    table = doc.add_table(rows=rows, cols=cols)
                    table.style = 'Table Grid'
                    
                    for i, row_data in enumerate(table_data["data"]):
                        for j, cell_data in enumerate(row_data):
                            if j < cols:
                                table.cell(i, j).text = cell_data["text"]
            
            # 문서 저장
            doc.save(output_path)
            logger.info(f"Word 템플릿 생성 완료: {output_path}")
            
            return output_path
            
        except Exception as e:
            logger.error(f"Word 템플릿 생성 실패: {e}")
            raise
    
    def convert_to_pdf_format(self, word_data: Dict[str, Any]) -> Dict[str, Any]:
        """Word 데이터를 PDF 형식으로 변환"""
        return {
            "pdf_info": {
                "total_pages": 1,  # Word는 페이지 개념이 다름
                "file_size": 0,
                "file_name": Path(word_data["file_path"]).name,
                "creation_date": 0
            },
            "toc_structure": {
                "sections": word_data["sections"],
                "chapters": [],
                "subsections": []
            },
            "split_results": [],
            "word_specific": {
                "tables": word_data["tables"],
                "images": word_data["images"],
                "categories": word_data["categories"]
            }
        }

# 사용 예시
def main():
    """Word 문서 처리 예시"""
    processor = WordDocumentProcessor()
    
    # 샘플 Word 문서 생성
    doc = Document()
    doc.add_heading('건설공사 내역서', 0)
    
    # 공통부문 섹션
    doc.add_heading('공통부문', 1)
    doc.add_paragraph('제1장 적용기준')
    doc.add_paragraph('1-1 일반사항')
    
    # 토목부문 섹션
    doc.add_heading('토목부문', 1)
    doc.add_paragraph('제2장 도로공사')
    doc.add_paragraph('2-1 도로포장')
    
    # 테이블 추가
    table = doc.add_table(rows=3, cols=3)
    table.style = 'Table Grid'
    table.cell(0, 0).text = '공종'
    table.cell(0, 1).text = '단위'
    table.cell(0, 2).text = '단가'
    
    # 문서 저장
    sample_path = Path("sample_document.docx")
    doc.save(sample_path)
    
    # 문서 분석
    if processor.can_process(sample_path):
        result = processor.extract_structured_data(sample_path)
        print("Word 문서 분석 결과:")
        print(json.dumps(result, ensure_ascii=False, indent=2))
    
    # 샘플 파일 정리
    sample_path.unlink()

if __name__ == "__main__":
    main() 